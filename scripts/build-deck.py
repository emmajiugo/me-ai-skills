"""Generate docs/me-ai-skills.pptx from the beats in docs/presentation.md.

Palette is index.html's dark mode, so the deck and the skill list read as one
project. Run: .cache/venv/bin/python scripts/build-deck.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG     = RGBColor(0x16, 0x15, 0x13)
PANEL  = RGBColor(0x1E, 0x1D, 0x1A)
INK    = RGBColor(0xEC, 0xE9, 0xE3)
MUTED  = RGBColor(0x97, 0x92, 0x8A)
LINE   = RGBColor(0x30, 0x2E, 0x2A)
ACCENT = RGBColor(0xD8, 0xA1, 0x5F)
GREEN  = RGBColor(0x7B, 0xB9, 0x8C)
RED    = RGBColor(0xC9, 0x7B, 0x6B)

SANS, MONO = "Helvetica Neue", "Menlo"
W, H = 13.333, 7.5
M = 0.95  # left margin

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)


def slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(s, x, y, w, h, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    return tf


def line(tf, txt, size, color=INK, font=SANS, bold=False, space_before=0,
         space_after=0, first=False, spacing=1.15):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before, p.space_after = Pt(space_before), Pt(space_after)
    p.line_spacing = spacing
    r = p.add_run()
    r.text = txt
    r.font.size, r.font.name, r.font.bold = Pt(size), font, bold
    r.font.color.rgb = color
    return p


def eyebrow(s, txt):
    tf = text(s, M, 0.72, 10, 0.4)
    line(tf, txt.upper(), 12, ACCENT, bold=True, first=True)


def headline(s, txt, y=1.25, size=40, w=11.4):
    tf = text(s, M, y, w, 2.2)
    line(tf, txt, size, INK, bold=True, first=True, spacing=1.05)


def panel(s, x, y, w, h, fill=PANEL, border=LINE):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.05
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = border
    box.line.width = Pt(1)
    box.shadow.inherit = False
    box.text_frame.text = ""
    return box


def code(s, x, y, w, lines, size=17, color=INK):
    h = 0.42 + len(lines) * (size * 1.6 / 72)
    panel(s, x, y, w, h)
    tf = text(s, x + 0.32, y + 0.21, w - 0.64, h - 0.42)
    for i, ln in enumerate(lines):
        c = MUTED if ln.startswith("#") else color
        line(tf, ln if ln else " ", size, c, font=MONO, first=(i == 0), spacing=1.35)
    return y + h


def bullets(s, items, y=3.0, size=21, gap=13, w=11.0, x=None):
    x = M if x is None else x
    tf = text(s, x, y, w, 4.0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else gap)
        p.line_spacing = 1.25
        d = p.add_run()
        d.text = "▪  "
        d.font.size, d.font.name = Pt(size), SANS
        d.font.color.rgb = ACCENT
        r = p.add_run()
        r.text = it
        r.font.size, r.font.name = Pt(size), SANS
        r.font.color.rgb = INK


def note(s, txt, y=6.55, color=MUTED, size=15, x=None, w=11.4):
    tf = text(s, M if x is None else x, y, w, 0.6)
    line(tf, txt, size, color, first=True, spacing=1.3)


def stat(s, x, y, w, h, label, big, big_color, sub, tail=None, tail_color=None):
    panel(s, x, y, w, h)
    tf = text(s, x + 0.42, y + 0.32, w - 0.84, 0.4)
    line(tf, label, 16, MUTED, font=MONO, first=True)
    tf = text(s, x + 0.42, y + 0.85, w - 0.84, 1.2)
    line(tf, big, 64, big_color, bold=True, first=True)
    tf = text(s, x + 0.42, y + 2.1, w - 0.84, 1.2)
    line(tf, sub, 19, INK, first=True, spacing=1.3)
    if tail:
        line(tf, tail, 18, tail_color or GREEN, space_before=10, spacing=1.3)


def deflist(s, rows, y, key_w=3.5, size=19, gap=0.72, x=None):
    x = M if x is None else x
    for i, (k, v) in enumerate(rows):
        yy = y + i * gap
        tf = text(s, x, yy, key_w, 0.5)
        line(tf, k, size, ACCENT, font=MONO, bold=True, first=True)
        tf = text(s, x + key_w + 0.25, yy + 0.02, 12.0 - key_w - M - 0.25, 0.6)
        line(tf, v, size - 2, MUTED, first=True, spacing=1.25)


# ---------------------------------------------------------------- 1. title
s = slide()
tf = text(s, M, 2.35, 11.4, 1.5)
line(tf, "43 skills to get one", 60, INK, bold=True, first=True)
tf = text(s, M, 3.75, 10.2, 1.8)
line(tf, "What Claude Code's skill curation already does,", 24, ACCENT, first=True, spacing=1.35)
line(tf, "where it runs out, and what I built past that point.", 24, ACCENT, spacing=1.35)
note(s, "Claude Code 2.1.267. Every number in this deck came off my own machine.", y=6.35)

# ------------------------------------------------------------ 2. I wanted one
s = slide()
eyebrow(s, "How this started")
headline(s, "I wanted one skill.", size=44)
tf = text(s, M, 3.0, 11.0, 2.0)
line(tf, "unslop strips the AI voice out of writing.", 24, INK, first=True, spacing=1.4)
line(tf, "It lives in cursor/plugins.", 24, MUTED, space_before=10, spacing=1.4)
line(tf, "Both of my sources are already Claude Code marketplaces. "
         "So I checked what it costs to take one skill out of each.", 21, MUTED,
     space_before=20, spacing=1.4)

# ------------------------------------------------------------- 3. the numbers
s = slide()
eyebrow(s, "What you actually install")
headline(s, "The plugin author picks the unit. Not you.", size=34)
cw = 5.35
stat(s, M, 2.75, cw, 3.5, "cursor/plugins → pstack", "43", ACCENT,
     "skills installed to get 1.", "22 of them are named principle-something.", MUTED)
stat(s, M + cw + 0.55, 2.75, cw, 3.5, "mattpocock/skills", "1", ACCENT,
     "plugin in the whole marketplace.", "21 skills. All or nothing.", MUTED)

# ----------------------------------------------------------- 4. what it costs
s = slide()
eyebrow(s, "The cost")
headline(s, "The context bill is the small problem.")
tf = text(s, M, 2.9, 11.2, 3.0)
line(tf, "Every skill Claude can see puts its description in context. "
         "Every session, every project.", 21, MUTED, first=True, spacing=1.4)
line(tf, "But descriptions are how the agent decides what to reach for.", 23, INK,
     space_before=22, spacing=1.4)
line(tf, "42 skills you did not want are 42 chances to pick the wrong one.",
     25, ACCENT, bold=True, space_before=14, spacing=1.4)
note(s, "You notice this as the agent doing something adjacent to what you asked. "
        "Then you blame the model.", y=6.3)

# ------------------------------------------------------------ 5. cant install
s = slide()
eyebrow(s, "The obvious answer")
headline(s, "\"So install fewer.\"\nYou can't.", size=40)
panel(s, M, 3.6, 11.0, 1.5)
tf = text(s, M + 0.5, 3.95, 10.0, 1.0)
line(tf, "\"The documentation does not provide a way to install individual skills "
         "from a plugin without installing the entire plugin.\"", 19, INK, first=True, spacing=1.35)
note(s, "Claude Code docs. Install granularity is the plugin.", y=5.4, size=17)

# --------------------------------------------------------- 6. skillOverrides
s = slide()
eyebrow(s, "What you do get")
headline(s, "skillOverrides", size=44)
note(s, "Per-skill visibility, set in settings.json, without editing the skill.",
     y=2.35, size=21, color=INK)
code(s, M, 3.05, 11.4, [
    "{",
    "  \"skillOverrides\": {",
    "    \"principle-laziness-protocol\": \"off\",",
    "    \"architect\": \"name-only\",",
    "    \"commit\": \"user-invocable-only\"",
    "  }",
    "}",
], size=18)

# ------------------------------------------------------------- 7. four levels
s = slide()
eyebrow(s, "What you do get")
headline(s, "Four levels.", size=42)
deflist(s, [
    ("on", "The default. Full description in context, shows in the slash menu."),
    ("name-only", "Claude sees the name, not the description. Available, but stops competing on trigger text."),
    ("user-invocable-only", "Hidden from Claude. You can still run it yourself."),
    ("off", "Hidden from Claude and from the slash menu."),
], y=2.85, key_w=4.15, size=20, gap=0.92)
note(s, "name-only is the one nobody finds by accident.", y=6.55, color=GREEN, size=19)

# --------------------------------------------------------- 8. rest of toolkit
s = slide()
eyebrow(s, "What you do get")
headline(s, "And the rest of the native toolkit.", size=36)
deflist(s, [
    ("disable-model-invocation", "In a skill's frontmatter. Makes it slash-only."),
    ("user-invocable: false", "The reverse. Hides it from your slash menu."),
    ("Skill(deploy *)", "Permission rules. Allow or deny by name."),
    ("disableBundledSkills", "Turns off everything bundled except /doctor."),
], y=2.75, key_w=5.0, size=19, gap=0.85)
note(s, "Then check it with /context. If off is working, the skill is not there.",
     y=6.4, color=GREEN, size=19)

# ------------------------------------------------------------ 9. honest turn
s = slide()
eyebrow(s, "Be straight about this")
headline(s, "So the volume problem is\nmostly solved already.", size=38)
tf = text(s, M, 3.8, 11.0, 2.2)
line(tf, "Install everything, turn off what you did not want. It works.", 23, INK,
     first=True, spacing=1.4)
line(tf, "If that is your whole problem: go set skillOverrides. You are done.",
     23, GREEN, bold=True, space_before=18, spacing=1.4)
line(tf, "The rest of this only matters once you have gone past it.", 21, MUTED,
     space_before=18, spacing=1.4)

# ----------------------------------------------------------- 10. four gaps
s = slide()
eyebrow(s, "Where it runs out")
headline(s, "Four things it does not do.", size=40)
bullets(s, [
    "Your edits do not survive updates.",
    "The list is opt-out, so it drifts. Skill 44 arrives enabled.",
    "It is Claude Code only. Plugins do nothing for opencode.",
    "Nothing records what you chose, only what you rejected.",
], y=2.9, size=22, gap=22, w=11.2)

# ------------------------------------------------------------- 11. the big one
s = slide()
eyebrow(s, "Where it runs out")
headline(s, "The first one is the real one.", size=38)
tf = text(s, M, 2.75, 11.2, 3.2)
line(tf, "skillOverrides controls visibility. Not content.", 23, INK, first=True, spacing=1.4)
line(tf, "The moment you want a skill to fire on your triggers instead of the "
         "author's, you edit their description inside the installed plugin.",
     21, MUTED, space_before=18, spacing=1.4)
line(tf, "The next update overwrites it. There is no supported place to put "
         "your version.", 23, RED, bold=True, space_before=18, spacing=1.4)
note(s, "That is a fork. You just did not decide to make one.", y=6.4, size=20, color=ACCENT)

# ------------------------------------------------------------- 12. the seam
s = slide()
eyebrow(s, "The idea worth stealing")
headline(s, "Put your changes and theirs where\nthey can never overlap.", y=1.15, size=34)
rows = [
    ("upstream",        "what the author wrote, untouched, committed", MUTED),
    ("your overrides",  "description, triggers, your own notes",       ACCENT),
    ("generated skill", "what the agent actually loads",               GREEN),
]
y = 3.02
for i, (a, b, c) in enumerate(rows):
    panel(s, M, y, 10.4, 0.82)
    tf = text(s, M + 0.42, y + 0.19, 3.6, 0.5)
    line(tf, a, 20, c, font=MONO, bold=True, first=True)
    tf = text(s, M + 4.3, y + 0.21, 5.8, 0.5)
    line(tf, b, 18, MUTED, first=True)
    if i < 2:
        tf = text(s, M + 0.42, y + 0.88, 1.0, 0.4)
        line(tf, "+" if i == 0 else "↓", 17, MUTED, first=True)
    y += 1.22
note(s, "No three-way merge anywhere, because there is nothing to merge. "
        "The price: you cannot edit the body, only add around it.", y=6.5, size=16)

# ------------------------------------------------------------ 13. composition
s = slide()
eyebrow(s, "The part I did not expect")
headline(s, "Once vendored, provenance\nstops mattering.", size=36)
note(s, "At runtime they are all just local skills that can call each other.", y=3.05, size=20)
code(s, M, 3.6, 8.2, [
    "to-tickets   →  mattpocock/skills",
    "unslop       →  cursor/plugins",
], size=19, color=GREEN)
note(s, "Two authors who have never heard of each other. Five lines of glue, and every "
        "ticket lands in Linear without the AI voice.", y=5.55, size=19)
note(s, "Do not concatenate two skills. Write a third that calls both. "
        "Then the build can link-check the chain.", y=6.6, size=16)

# ------------------------------------------------------------ 14. when not to
s = slide()
eyebrow(s, "Be honest")
headline(s, "When not to use this.", size=42)
bullets(s, [
    "Your problem is volume. skillOverrides is enough, and this is more machinery than you need.",
    "You never rewrite descriptions. You will never hit the seam.",
    "You want to change what a skill does, not when it fires. You will eject immediately.",
    "Your team needs one shared set. This installs globally, per developer.",
], y=2.75, size=20, gap=18, w=11.2)

# ---------------------------------------------------------------- 15. close
s = slide()
eyebrow(s, "The one line")
headline(s, "Decide where the seam goes\nbefore you make the edit.", y=2.15, size=40)
tf = text(s, M, 4.6, 11.2, 1.8)
line(tf, "Otherwise you have forked. You will find out months later, "
         "looking for a fix that was published and never reached you.",
     22, MUTED, first=True, spacing=1.4)

prs.save("docs/me-ai-skills.pptx")
print(f"docs/me-ai-skills.pptx  {len(prs.slides._sldIdLst)} slides")
